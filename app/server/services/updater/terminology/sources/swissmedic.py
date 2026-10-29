# Copyright © 2023–2025 Thomas Virdis
# Licensed under the GNU General Public License, version 3 or later.

"""Swissmedic list of authorised human medicinal product packs (open use).

Each dosage strength becomes a Swiss product. Active ingredients are given as
Latin pharmacopoeial names; they are resolved against existing verified
terminology by deterministic Latin-to-English variants, never by fuzzy
matching. The source creates no ingredient concepts of its own.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import openpyxl

from services.terminology.normalization import fold, latin_to_english_variants
from services.updater.terminology.download import download_file
from services.updater.terminology.snapshot import (
    ProgressFn,
    StagedProduct,
    StagedProductIngredient,
    StopFn,
    TerminologySnapshot,
    TerminologySourceUnavailable,
)

PACKS_URL = (
    "https://www.swissmedic.ch/dam/swissmedic/en/dokumente/internetlisten/"
    "zugelassene_packungen_human.xlsx.download.xlsx/zugelassene_packungen_ham.xlsx"
)
LICENSE = "Swissmedic open data (open use; source attribution required)"
ATTRIBUTION = "Swissmedic, Swiss Agency for Therapeutic Products"
HEADER_PREFIXES = {
    "authorisation": "zulassungs",
    "dosage": "dosisst",
    "name": "bezeichnung",
    "atc": "atc",
    "ingredients": "wirkstoff",
    "composition": "zusammensetzung",
}
_STAND_RE = re.compile(r"(\d{2})\.(\d{2})\.(\d{4})")
_BRAND_NOISE_RE = re.compile(r"\([^)]*\)|\b\d+(?:[.,']\d+)*\s*(?:mg|g|µg|μg|mcg|ml|%|i\.e\.|u\.i\.)?(?=\s|$)", re.IGNORECASE)
_STRENGTH_TEMPLATE = r"\s+(\d+(?:[.,]\d+)?)\s*(mg|g|µg|μg|mcg|ml|u\.i\.|%)(?![a-z])"

###############################################################################
class SwissmedicSource:
    key = "swissmedic"
    license = LICENSE
    attribution = ATTRIBUTION
    precedence = 4

    # -------------------------------------------------------------------------
    def __init__(self, archive_dir: Path, *, timeout: float = 300.0) -> None:
        self.archive_dir = archive_dir
        self.timeout = timeout

    # -------------------------------------------------------------------------
    def probe_version(self) -> str | None:
        # Swissmedic publishes no ETag or Last-Modified; the content hash
        # computed after download decides whether the snapshot changed.
        return None

    # -------------------------------------------------------------------------
    def fetch(self, progress: ProgressFn, should_stop: StopFn) -> TerminologySnapshot:
        progress(5.0, "Downloading Swissmedic authorised packs")
        workbook, sha256 = download_file(
            PACKS_URL,
            self.archive_dir / "swissmedic_packs_human.xlsx",
            timeout=self.timeout,
            should_stop=should_stop,
        )
        progress(40.0, "Parsing Swissmedic products")
        return self.parse_workbook(workbook, content_hash=sha256, should_stop=should_stop)

    # -------------------------------------------------------------------------
    def parse_workbook(
        self,
        path: Path,
        *,
        content_hash: str,
        should_stop: StopFn,
    ) -> TerminologySnapshot:
        try:
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 - openpyxl raises several types
            raise TerminologySourceUnavailable(f"Swissmedic workbook unreadable: {exc}") from exc
        try:
            sheet = workbook[workbook.sheetnames[0]]
            rows = sheet.iter_rows(values_only=True)
            columns: dict[str, int] | None = None
            stand: str | None = None
            products: dict[str, StagedProduct] = {}
            for row in rows:
                if columns is None:
                    stand = stand or self._find_stand(row)
                    columns = self._detect_header(row)
                    continue
                if should_stop():
                    raise TerminologySourceUnavailable("Swissmedic parse cancelled")
                product = self._parse_row(row, columns)
                if product is not None and product.product_id not in products:
                    products[product.product_id] = product
        finally:
            workbook.close()
        if columns is None:
            raise TerminologySourceUnavailable("Swissmedic workbook header not found")
        version = f"{stand or 'unknown'} sha256:{content_hash[:12]}"
        snapshot = TerminologySnapshot(
            source=self.key,
            version=version,
            license=self.license,
            products=list(products.values()),
            brand_families=True,
        )
        snapshot.stats.update(products=len(snapshot.products))
        return snapshot

    # -------------------------------------------------------------------------
    @staticmethod
    def _find_stand(row: tuple[Any, ...]) -> str | None:
        for cell in row:
            match = _STAND_RE.search(str(cell or ""))
            if match and "stand" in str(cell).casefold():
                day, month, year = match.groups()
                return f"{year}-{month}-{day}"
        return None

    # -------------------------------------------------------------------------
    @staticmethod
    def _detect_header(row: tuple[Any, ...]) -> dict[str, int] | None:
        found: dict[str, int] = {}
        for index, cell in enumerate(row):
            label = str(cell or "").strip().split("\n", 1)[0].casefold()
            for column, prefix in HEADER_PREFIXES.items():
                if column not in found and label.startswith(prefix):
                    found[column] = index
        return found if len(found) == len(HEADER_PREFIXES) else None

    # -------------------------------------------------------------------------
    def _parse_row(self, row: tuple[Any, ...], columns: dict[str, int]) -> StagedProduct | None:
        def cell(name: str) -> str:
            index = columns[name]
            return str(row[index]).strip() if index < len(row) and row[index] is not None else ""

        authorisation, dosage, name = cell("authorisation"), cell("dosage"), cell("name")
        actives = cell("ingredients")
        if not authorisation or not name or not actives:
            return None
        composition = cell("composition")
        ingredients: list[StagedProductIngredient] = []
        for active in (part.strip() for part in actives.split(",")):
            if not active:
                continue
            folded = fold(active)
            strength = re.search(re.escape(active) + _STRENGTH_TEMPLATE, composition, re.IGNORECASE)
            ingredients.append(
                StagedProductIngredient(
                    ingredient_name=active,
                    lookup_names=tuple(dict.fromkeys([folded, *latin_to_english_variants(folded)])),
                    strength_value=float(strength.group(1).replace(",", ".")) if strength else None,
                    strength_unit=strength.group(2).casefold() if strength else None,
                )
            )
        if not ingredients:
            return None
        head, _, form = name.partition(", ")
        return StagedProduct(
            product_id=f"{authorisation}-{dosage or '0'}",
            name=name,
            brand_name=self.brand_name(head),
            ingredients=ingredients,
            country="CH",
            dose_form=form.strip() or None,
            atc_code=cell("atc") or None,
        )

    # -------------------------------------------------------------------------
    @staticmethod
    def brand_name(head: str) -> str | None:
        cleaned = " ".join(_BRAND_NOISE_RE.sub(" ", head).split())
        return cleaned or None
